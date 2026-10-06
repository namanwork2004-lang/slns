"""Shree Laxmi Narayan Steels - production tracking app.

Lots flow Inward -> 18 G Dhulai -> Machine -> Tapai -> Final -> Complete.
Every movement is stored as (lot, stage, weight, date). Rules enforced here:
  * a lot's weight never changes between stages;
  * by any date, the weight that has reached a stage can never exceed the
    weight that had reached the previous stage (Inward is capped by lot weight),
    so the same material can't pass through a stage twice or skip a stage;
  * partial moves are allowed.

Storage: Postgres when DATABASE_URL (or POSTGRES_URL) is set, e.g. on Vercel
with Neon; otherwise a local SQLite file, factory.db.
Login: one shared password from the APP_PASSWORD environment variable.
"""
import csv
import datetime as dt
import hashlib
import hmac
import io
import os
import sqlite3

from flask import Flask, g, jsonify, redirect, request, session

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE, "factory.db")
DATABASE_URL = os.environ.get("DATABASE_URL") or os.environ.get("POSTGRES_URL")
APP_PASSWORD = os.environ.get("APP_PASSWORD", "")
ON_VERCEL = bool(os.environ.get("VERCEL"))
IST = dt.timezone(dt.timedelta(hours=5, minutes=30))
EPS = 0.05  # kg tolerance for decimal rounding

STAGES = [
    {"key": "inward", "name": "Inward", "hi": "आवक माल"},
    {"key": "dhulai", "name": "18 G Dhulai", "hi": "18 G धुलाई"},
    {"key": "machine", "name": "Machine", "hi": "मशीन"},
    {"key": "tapai", "name": "Tapai", "hi": "तपाई (भट्टी)"},
    {"key": "final", "name": "Final", "hi": "फाइनल माल"},
    {"key": "complete", "name": "Complete", "hi": "तैयार माल (Ready)"},
]

app = Flask(__name__, static_folder="public", static_url_path="")
# Changing APP_PASSWORD changes the key, which signs everyone out.
app.secret_key = hashlib.sha256(f"slns|{APP_PASSWORD}|{DATABASE_URL}".encode()).hexdigest()
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
                  SESSION_COOKIE_SECURE=ON_VERCEL,
                  PERMANENT_SESSION_LIFETIME=dt.timedelta(days=30))


def today():
    return dt.datetime.now(IST).date()


# ---------------------------------------------------------------- database

class Database:
    """Thin wrapper so the same SQL (with ? placeholders) runs on SQLite and Postgres."""

    def __init__(self):
        if DATABASE_URL:
            import psycopg
            from psycopg.rows import dict_row
            self.con = psycopg.connect(DATABASE_URL, row_factory=dict_row)
        else:
            self.con = sqlite3.connect(DB_PATH)
            self.con.row_factory = sqlite3.Row
            self.con.execute("PRAGMA foreign_keys=ON")

    def q(self, sql, args=()):
        if DATABASE_URL:
            sql = sql.replace("?", "%s")
        return self.con.execute(sql, args)

    def many(self, sql, rows):
        if DATABASE_URL:
            with self.con.cursor() as cur:
                cur.executemany(sql.replace("?", "%s"), rows)
        else:
            self.con.executemany(sql, rows)

    def commit(self):
        self.con.commit()

    def rollback(self):
        self.con.rollback()

    def close(self):
        self.con.close()


SCHEMA = """
CREATE TABLE IF NOT EXISTS lots(
    lot_no INTEGER PRIMARY KEY,
    qty TEXT NOT NULL DEFAULT '',
    size TEXT NOT NULL DEFAULT '',
    pc TEXT NOT NULL DEFAULT '',
    wt DOUBLE PRECISION NOT NULL CHECK (wt > 0));
CREATE TABLE IF NOT EXISTS moves(
    id {id_col},
    lot_no INTEGER NOT NULL REFERENCES lots(lot_no),
    stage INTEGER NOT NULL,
    wt DOUBLE PRECISION NOT NULL CHECK (wt > 0),
    date TEXT NOT NULL,
    created_at TEXT NOT NULL);
CREATE INDEX IF NOT EXISTS moves_lot ON moves(lot_no);
"""
_schema_ready = False


def db():
    global _schema_ready
    if "db" not in g:
        g.db = Database()
        if not _schema_ready:
            id_col = "SERIAL PRIMARY KEY" if DATABASE_URL else "INTEGER PRIMARY KEY AUTOINCREMENT"
            for stmt in SCHEMA.format(id_col=id_col).split(";"):
                if stmt.strip():
                    g.db.q(stmt)
            g.db.commit()
            _schema_ready = True
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    con = g.pop("db", None)
    if con is not None:
        con.close()


def parse_rows(rows):
    """Turn spreadsheet rows (header LOT, QTY, SIZE, PC, WT) into lot tuples."""
    header_idx = None
    for i, r in enumerate(rows):
        cells = [str(c if c is not None else "").strip().upper() for c in r]
        if "LOT" in cells and "WT" in cells:
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("Could not find a header row with LOT and WT columns.")
    cells = [str(c if c is not None else "").strip().upper() for c in rows[header_idx]]
    col = {k: cells.index(k) for k in ("LOT", "QTY", "SIZE", "PC", "WT") if k in cells}

    def cell(r, k):
        if k not in col or col[k] >= len(r) or r[col[k]] is None:
            return ""
        v = r[col[k]]
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).strip()

    out, seen = [], set()
    for n, r in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        lot, wt = cell(r, "LOT"), cell(r, "WT")
        if not lot and not wt:
            continue
        try:
            lot_no, w = int(float(lot)), float(wt)
        except ValueError:
            raise ValueError(f"Row {n}: LOT '{lot}' / WT '{wt}' is not a number.")
        if w <= 0:
            raise ValueError(f"Row {n}: weight for lot {lot_no} must be more than 0.")
        if lot_no in seen:
            raise ValueError(f"Row {n}: lot {lot_no} appears more than once.")
        seen.add(lot_no)
        out.append((lot_no, cell(r, "QTY"), cell(r, "SIZE"), cell(r, "PC"), round(w, 3)))
    return out


# ------------------------------------------------------------------ rules

def fmt_date(iso):
    return dt.date.fromisoformat(iso).strftime("%d-%m-%Y")


def check_lot(con, lot_no):
    """Return an error message if this lot's movements break the stage rules."""
    row = con.q("SELECT wt FROM lots WHERE lot_no=?", (lot_no,)).fetchone()
    if row is None:
        return f"Lot {lot_no} does not exist."
    wt = row["wt"]
    moves = con.q("SELECT stage, wt, date FROM moves WHERE lot_no=?", (lot_no,)).fetchall()
    for day in sorted({m["date"] for m in moves}):
        cum = [0.0] * len(STAGES)
        for m in moves:
            if m["date"] <= day:
                cum[m["stage"]] += m["wt"]
        if cum[0] > wt + EPS:
            return (f"Lot {lot_no}: Inward would be {cum[0]:,.1f} kg, "
                    f"more than the lot weight of {wt:,.1f} kg.")
        for s in range(1, len(STAGES)):
            if cum[s] > cum[s - 1] + EPS:
                return (f"Lot {lot_no}: by {fmt_date(day)}, {STAGES[s]['name']} would have "
                        f"{cum[s]:,.1f} kg but only {cum[s - 1]:,.1f} kg had reached "
                        f"{STAGES[s - 1]['name']}.")
    return None


def valid_date(value):
    try:
        d = dt.date.fromisoformat(value)
    except (TypeError, ValueError):
        return None
    return d if d <= today() else None


def bad(msg, code=400):
    return jsonify({"error": msg}), code


UPSERT_LOT = """INSERT INTO lots(lot_no, qty, size, pc, wt) VALUES (?,?,?,?,?)
    ON CONFLICT(lot_no) DO UPDATE SET qty=excluded.qty, size=excluded.size,
    pc=excluded.pc, wt=excluded.wt"""


# ------------------------------------------------------------------- auth

@app.before_request
def require_login():
    if not request.path.startswith("/api/"):
        return None
    if ON_VERCEL and not DATABASE_URL:
        return bad("Database not connected. In Vercel, add a Neon (Postgres) database "
                   "under Storage, connect it to this project, then redeploy.", 503)
    if request.path == "/api/login":
        return None
    if not APP_PASSWORD:
        if ON_VERCEL:
            return bad("APP_PASSWORD is not set on the server.", 503)
        return None  # local testing without a password
    if not session.get("ok"):
        return bad("Please log in.", 401)
    return None


@app.post("/api/login")
def login():
    pw = str((request.get_json(silent=True) or {}).get("password", ""))
    if not APP_PASSWORD or not hmac.compare_digest(pw.encode(), APP_PASSWORD.encode()):
        return bad("Wrong password.", 401)
    session.permanent = True
    session["ok"] = True
    return jsonify({"ok": True})


@app.post("/api/logout")
def logout():
    session.clear()
    return jsonify({"ok": True})


# -------------------------------------------------------------------- api

@app.get("/")
def index():
    # On Vercel, files in public/ are served by the CDN, not bundled with this function.
    if os.path.exists(os.path.join(BASE, "public", "index.html")):
        return app.send_static_file("index.html")
    return redirect("/index.html")


@app.get("/api/state")
def state():
    con = db()
    lots = [dict(r) for r in con.q("SELECT * FROM lots ORDER BY lot_no").fetchall()]
    moves = [dict(r) for r in con.q("SELECT * FROM moves ORDER BY date, id").fetchall()]
    return jsonify({"stages": STAGES, "lots": lots, "moves": moves,
                    "today": today().isoformat(), "auth": bool(APP_PASSWORD)})


@app.post("/api/moves")
def add_moves():
    body = request.get_json(silent=True) or {}
    stage, items = body.get("stage"), body.get("items") or []
    if not isinstance(stage, int) or not 0 <= stage < len(STAGES):
        return bad("Unknown stage.")
    day = valid_date(body.get("date"))
    if day is None:
        return bad("Pick a valid date that is not in the future.")
    if not items:
        return bad("Select at least one lot.")
    con = db()
    now = dt.datetime.now(IST).isoformat(timespec="seconds")
    lots = set()
    try:
        for it in items:
            lot_no, wt = int(it["lot_no"]), round(float(it["wt"]), 3)
            if wt <= 0:
                raise ValueError(f"Lot {lot_no}: weight must be more than 0.")
            if con.q("SELECT 1 FROM lots WHERE lot_no=?", (lot_no,)).fetchone() is None:
                raise ValueError(f"Lot {lot_no} does not exist.")
            con.q("INSERT INTO moves(lot_no, stage, wt, date, created_at) VALUES (?,?,?,?,?)",
                  (lot_no, stage, wt, day.isoformat(), now))
            lots.add(lot_no)
        errors = [e for e in (check_lot(con, n) for n in sorted(lots)) if e]
        if errors:
            con.rollback()
            return bad(" ".join(errors))
    except (KeyError, TypeError, ValueError) as e:
        con.rollback()
        return bad(str(e) if isinstance(e, ValueError) else "Invalid lot or weight.")
    con.commit()
    return jsonify({"ok": True, "count": len(items)})


@app.put("/api/moves/<int:move_id>")
def edit_move(move_id):
    b = request.get_json(silent=True) or {}
    con = db()
    old = con.q("SELECT lot_no FROM moves WHERE id=?", (move_id,)).fetchone()
    if old is None:
        return bad("Entry not found.", 404)
    try:
        lot_no, stage, wt = int(b["lot_no"]), int(b["stage"]), round(float(b["wt"]), 3)
    except (KeyError, TypeError, ValueError):
        return bad("Lot, stage and weight must be numbers.")
    if not 0 <= stage < len(STAGES):
        return bad("Unknown stage.")
    if wt <= 0:
        return bad("Weight must be more than 0.")
    day = valid_date(b.get("date"))
    if day is None:
        return bad("Pick a valid date that is not in the future.")
    if con.q("SELECT 1 FROM lots WHERE lot_no=?", (lot_no,)).fetchone() is None:
        return bad(f"Lot {lot_no} does not exist.")
    con.q("UPDATE moves SET lot_no=?, stage=?, wt=?, date=? WHERE id=?",
          (lot_no, stage, wt, day.isoformat(), move_id))
    errors = [e for e in (check_lot(con, n) for n in sorted({old["lot_no"], lot_no})) if e]
    if errors:
        con.rollback()
        return bad("Can't save this change. " + " ".join(errors))
    con.commit()
    return jsonify({"ok": True})


@app.delete("/api/moves/<int:move_id>")
def undo_move(move_id):
    con = db()
    row = con.q("SELECT lot_no FROM moves WHERE id=?", (move_id,)).fetchone()
    if row is None:
        return bad("Entry not found.", 404)
    con.q("DELETE FROM moves WHERE id=?", (move_id,))
    err = check_lot(con, row["lot_no"])
    if err:
        con.rollback()
        return bad("Can't undo: this material has already moved to the next stage. "
                   "Undo that entry first. (" + err + ")")
    con.commit()
    return jsonify({"ok": True})


@app.post("/api/lots")
def save_lot():
    b = request.get_json(silent=True) or {}
    try:
        lot_no, wt = int(b["lot_no"]), round(float(b["wt"]), 3)
    except (KeyError, TypeError, ValueError):
        return bad("Lot number and weight must be numbers.")
    if lot_no <= 0 or wt <= 0:
        return bad("Lot number and weight must be more than 0.")
    con = db()
    con.q(UPSERT_LOT, (lot_no, str(b.get("qty", "")).strip(), str(b.get("size", "")).strip(),
                       str(b.get("pc", "")).strip(), wt))
    err = check_lot(con, lot_no)
    if err:
        con.rollback()
        return bad(err)
    con.commit()
    return jsonify({"ok": True})


@app.delete("/api/lots/<int:lot_no>")
def delete_lot(lot_no):
    con = db()
    if con.q("SELECT 1 FROM moves WHERE lot_no=? LIMIT 1", (lot_no,)).fetchone():
        return bad(f"Lot {lot_no} already has production entries, so it can't be deleted.")
    con.q("DELETE FROM lots WHERE lot_no=?", (lot_no,))
    con.commit()
    return jsonify({"ok": True})


@app.post("/api/lots/import")
def import_lots():
    f = request.files.get("file")
    if f is None or not f.filename:
        return bad("Choose an Excel (.xlsx) or CSV file.")
    name = f.filename.lower()
    try:
        if name.endswith(".xlsx"):
            import openpyxl
            ws = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True).active
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
        elif name.endswith(".csv"):
            rows = list(csv.reader(io.StringIO(f.read().decode("utf-8-sig"))))
        else:
            return bad("Only .xlsx and .csv files are supported.")
        parsed = parse_rows(rows)
    except ValueError as e:
        return bad(str(e))
    except Exception:
        return bad("Could not read that file.")
    if not parsed:
        return bad("No lots found in that file.")
    con = db()
    existing = {r["lot_no"] for r in con.q("SELECT lot_no FROM lots").fetchall()}
    con.many(UPSERT_LOT, parsed)
    errors = [e for e in (check_lot(con, p[0]) for p in parsed if p[0] in existing) if e]
    if errors:
        con.rollback()
        return bad(" ".join(errors[:5]))
    con.commit()
    added = sum(1 for p in parsed if p[0] not in existing)
    return jsonify({"ok": True, "added": added, "updated": len(parsed) - added})


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", 5055)), debug=False)
